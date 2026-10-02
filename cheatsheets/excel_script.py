import formulas
import pandas as pd
import streamlit as st
from openpyxl.utils.cell import get_column_letter, range_boundaries

from .paths import DATA_DIR


def resolve_formula_reference(reference, cell_values):
    cell_range = reference.rsplit("!", 1)[-1].replace("$", "")
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    values = [
        [
            cell_values.get(f"{get_column_letter(col)}{row}", 0)
            for col in range(min_col, max_col + 1)
        ]
        for row in range(min_row, max_row + 1)
    ]
    if min_col == max_col and min_row == max_row:
        return values[0][0]
    return values


def calculate_formula(expression, cell_values):
    compiled = formulas.Parser().ast(expression)[1].compile()
    inputs = {
        reference: resolve_formula_reference(reference, cell_values)
        for reference in compiled.inputs
    }
    result = compiled(**inputs)
    if hasattr(result, "tolist"):
        result = result.tolist()
    return result


def excel_cheatsheet():
    st.set_page_config(page_title="CheatSheets | Excel", page_icon=":material/auto_stories:", layout="wide")
    workbook_file = DATA_DIR / "excel_cheatsheet.xlsx"
    if not workbook_file.is_file():
        st.error(f"Workbook not found: {workbook_file.name}")
        st.stop()

    st.title("Excel cheatsheet")
    st.download_button(
        "Download Excel workbook",
        data=workbook_file.read_bytes(),
        file_name=workbook_file.name,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        icon=":material/download:",
    )

    cheatsheet_tab, playground_tab = st.tabs(["Cheatsheet", "Playground"])
    with cheatsheet_tab:
        cheatsheet = pd.read_excel(workbook_file, sheet_name="Formula Cheatsheet")
        st.dataframe(cheatsheet, hide_index=True, width="stretch")

    with playground_tab:
        st.caption("Edit the sample cells, then calculate a formula using references such as A1 or A1:A5.")
        practice_grid = pd.read_excel(workbook_file, sheet_name="Playground Data", header=None)
        practice_grid.columns = [get_column_letter(index) for index in range(1, len(practice_grid.columns) + 1)]
        edited_grid = st.data_editor(
            practice_grid,
            hide_index=True,
            num_rows="dynamic",
            width="stretch",
            key="excel_playground_grid",
        )
        formula = st.text_input(
            "Excel formula",
            value="=SUM(A1:A5)+SUM(B1:B5)",
            key="excel_playground_formula",
        )

        cell_values = {}
        for row_index, row in edited_grid.iterrows():
            for column_index, value in enumerate(row, start=1):
                if pd.isna(value):
                    value = 0
                address = f"{get_column_letter(column_index)}{row_index + 1}"
                cell_values[address] = value

        if st.button("Calculate", icon=":material/functions:", key="calculate_excel_formula"):
            try:
                result = calculate_formula(formula, cell_values)
                st.session_state["excel_formula_result"] = result
                st.session_state.pop("excel_formula_error", None)
            except Exception as error:
                st.session_state.pop("excel_formula_result", None)
                st.session_state["excel_formula_error"] = str(error)

        if "excel_formula_result" in st.session_state:
            st.success(f"Result: {st.session_state['excel_formula_result']}")
        if st.session_state.get("excel_formula_error"):
            st.error(st.session_state["excel_formula_error"])



